"""
Piece 10 — Integration & Edge Case Hardening

Tests:
- Scanned PDF: pipeline returns warning, sets is_scanned, does not crash
- Password-protected PDF: pipeline skips file, returns clear message
- LLM unreachable: pipeline falls back to regex-only, sets llm_warning flag
- LLM returns garbage: falls back to regex-only, sets llm_warning flag
- Huge document: chunked correctly, mappings merged (no dupes from overlap)
- Mixed batch: PDF and DOCX processed in same job, each by correct reader
- Overlapping PII: name and email in same sentence both replaced independently
- PII in table cells: replaced in DOCX table output
- All 8 edge cases from the spec exercise their code paths
"""

import pytest
from conftest import KNOWN_PII, SAMPLE_PII_TEXT


# ---------------------------------------------------------------------------
# Helper: run full anonymize pipeline via the test client
# ---------------------------------------------------------------------------

def anonymize_file(app_client, file_path, mime_type, mock_fixture=None):
    with open(file_path, "rb") as f:
        upload = app_client.post(
            "/upload",
            files={"files": (file_path.name, f, mime_type)},
        )
    assert upload.status_code == 200
    job_id = upload.json()["job_id"]
    result = app_client.post(f"/jobs/{job_id}/anonymize")
    return job_id, result


PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ---------------------------------------------------------------------------
# Scanned PDF guard
# ---------------------------------------------------------------------------

class TestScannedPDFGuard:
    def test_scanned_pdf_does_not_crash(self, app_client, scanned_pdf_path, mock_llm_endpoint):
        job_id, result = anonymize_file(app_client, scanned_pdf_path, PDF_MIME)
        assert result.status_code in (200, 422)

    def test_scanned_pdf_returns_warning(self, app_client, scanned_pdf_path, mock_llm_endpoint):
        job_id, result = anonymize_file(app_client, scanned_pdf_path, PDF_MIME)
        data = result.json()
        warnings = data.get("warnings", [])
        assert any("scan" in w.lower() or "ocr" in w.lower() for w in warnings), (
            "Expected a warning about scanned PDF, got: " + str(warnings)
        )

    def test_scanned_pdf_job_status_not_complete(self, app_client, scanned_pdf_path, mock_llm_endpoint):
        job_id, _ = anonymize_file(app_client, scanned_pdf_path, PDF_MIME)
        jobs = app_client.get("/jobs").json()
        job = next(j for j in jobs if j["id"] == job_id)
        assert job["status"] in ("error", "skipped", "complete_with_warnings")

    def test_scanned_pdf_file_flagged_in_result(self, app_client, scanned_pdf_path, mock_llm_endpoint):
        job_id, result = anonymize_file(app_client, scanned_pdf_path, PDF_MIME)
        data = result.json()
        for file_result in data.get("files", []):
            if file_result["filename"] == scanned_pdf_path.name:
                assert file_result.get("is_scanned") is True


# ---------------------------------------------------------------------------
# Password-protected file guard
# ---------------------------------------------------------------------------

class TestPasswordProtectedGuard:
    def test_password_pdf_does_not_crash(self, app_client, password_protected_pdf_path, mock_llm_endpoint):
        job_id, result = anonymize_file(app_client, password_protected_pdf_path, PDF_MIME)
        assert result.status_code in (200, 422)

    def test_password_pdf_returns_warning(self, app_client, password_protected_pdf_path, mock_llm_endpoint):
        job_id, result = anonymize_file(app_client, password_protected_pdf_path, PDF_MIME)
        data = result.json()
        warnings = data.get("warnings", [])
        assert any("password" in w.lower() or "protect" in w.lower() for w in warnings)

    def test_password_pdf_flagged_in_result(self, app_client, password_protected_pdf_path, mock_llm_endpoint):
        job_id, result = anonymize_file(app_client, password_protected_pdf_path, PDF_MIME)
        data = result.json()
        for file_result in data.get("files", []):
            if file_result["filename"] == password_protected_pdf_path.name:
                assert file_result.get("is_password_protected") is True


# ---------------------------------------------------------------------------
# LLM unreachable — fallback to regex-only
# ---------------------------------------------------------------------------

class TestLLMUnreachableFallback:
    def test_pipeline_does_not_crash_when_llm_down(
        self, app_client, sample_pdf_path, mock_llm_unreachable
    ):
        job_id, result = anonymize_file(app_client, sample_pdf_path, PDF_MIME)
        assert result.status_code == 200

    def test_pipeline_returns_llm_warning(
        self, app_client, sample_pdf_path, mock_llm_unreachable
    ):
        job_id, result = anonymize_file(app_client, sample_pdf_path, PDF_MIME)
        data = result.json()
        warnings = data.get("warnings", [])
        assert any(
            "llm" in w.lower() or "ollama" in w.lower() or "fallback" in w.lower()
            for w in warnings
        )

    def test_regex_findings_still_applied_when_llm_down(
        self, app_client, sample_pdf_path, mock_llm_unreachable
    ):
        job_id, result = anonymize_file(app_client, sample_pdf_path, PDF_MIME)
        data = result.json()
        # At minimum, SSN should be replaced via regex
        for file_result in data.get("files", []):
            anonymized = file_result.get("anonymized_text", "")
            assert KNOWN_PII["ssn"] not in anonymized, (
                "SSN should be replaced by regex even when LLM is down"
            )


# ---------------------------------------------------------------------------
# LLM returns garbage — fallback to regex-only
# ---------------------------------------------------------------------------

class TestLLMGarbageFallback:
    def test_pipeline_does_not_crash_on_garbage_llm(
        self, app_client, sample_pdf_path, mock_llm_garbage
    ):
        job_id, result = anonymize_file(app_client, sample_pdf_path, PDF_MIME)
        assert result.status_code == 200

    def test_garbage_llm_sets_warning_flag(
        self, app_client, sample_pdf_path, mock_llm_garbage
    ):
        job_id, result = anonymize_file(app_client, sample_pdf_path, PDF_MIME)
        data = result.json()
        warnings = data.get("warnings", [])
        assert any(
            "fallback" in w.lower() or "llm" in w.lower() or "warning" in w.lower()
            for w in warnings
        )

    def test_regex_ssn_still_replaced_on_garbage_llm(
        self, app_client, sample_pdf_path, mock_llm_garbage
    ):
        job_id, result = anonymize_file(app_client, sample_pdf_path, PDF_MIME)
        data = result.json()
        for file_result in data.get("files", []):
            assert KNOWN_PII["ssn"] not in file_result.get("anonymized_text", "")


# ---------------------------------------------------------------------------
# Huge document chunking
# ---------------------------------------------------------------------------

class TestHugeDocumentChunking:
    def test_large_docx_anonymized_without_crash(
        self, app_client, multi_page_docx_path, mock_llm_endpoint
    ):
        job_id, result = anonymize_file(app_client, multi_page_docx_path, DOCX_MIME)
        assert result.status_code == 200

    def test_large_docx_pii_replaced(
        self, app_client, multi_page_docx_path, mock_llm_endpoint
    ):
        job_id, result = anonymize_file(app_client, multi_page_docx_path, DOCX_MIME)
        data = result.json()
        for file_result in data.get("files", []):
            anonymized = file_result.get("anonymized_text", "")
            assert KNOWN_PII["ssn"] not in anonymized

    def test_large_doc_mapping_has_no_duplicates(
        self, app_client, multi_page_docx_path, mock_llm_endpoint
    ):
        job_id, result = anonymize_file(app_client, multi_page_docx_path, DOCX_MIME)
        mapping = app_client.get(f"/jobs/{job_id}/mapping").json()
        originals = [e["original"] for e in mapping]
        # No original value should appear twice
        assert len(originals) == len(set(originals)), (
            "Duplicate entries in mapping from chunk overlap"
        )

    def test_large_doc_same_pii_same_placeholder_throughout(
        self, app_client, multi_page_docx_path, mock_llm_endpoint
    ):
        """The SSN appears in every paragraph — it must always map to [SSN_1]."""
        job_id, result = anonymize_file(app_client, multi_page_docx_path, DOCX_MIME)
        data = result.json()
        for file_result in data.get("files", []):
            anonymized = file_result.get("anonymized_text", "")
            # Only one SSN placeholder should exist
            import re
            ssn_placeholders = set(re.findall(r"\[SSN_\d+\]", anonymized))
            assert len(ssn_placeholders) <= 1, (
                f"Multiple SSN placeholders: {ssn_placeholders}"
            )


# ---------------------------------------------------------------------------
# Mixed file batch
# ---------------------------------------------------------------------------

class TestMixedFileBatch:
    def test_pdf_and_docx_in_same_job(
        self, app_client, sample_pdf_path, sample_docx_path, mock_llm_endpoint
    ):
        with open(sample_pdf_path, "rb") as f1, open(sample_docx_path, "rb") as f2:
            upload = app_client.post(
                "/upload",
                files=[
                    ("files", (sample_pdf_path.name, f1, PDF_MIME)),
                    ("files", (sample_docx_path.name, f2, DOCX_MIME)),
                ],
            )
        job_id = upload.json()["job_id"]
        result = app_client.post(f"/jobs/{job_id}/anonymize")
        assert result.status_code == 200
        data = result.json()
        filenames = [f["filename"] for f in data.get("files", [])]
        assert sample_pdf_path.name in filenames
        assert sample_docx_path.name in filenames

    def test_each_file_uses_correct_reader(
        self, app_client, sample_pdf_path, sample_docx_path, mock_llm_endpoint
    ):
        with open(sample_pdf_path, "rb") as f1, open(sample_docx_path, "rb") as f2:
            upload = app_client.post(
                "/upload",
                files=[
                    ("files", (sample_pdf_path.name, f1, PDF_MIME)),
                    ("files", (sample_docx_path.name, f2, DOCX_MIME)),
                ],
            )
        job_id = upload.json()["job_id"]
        result = app_client.post(f"/jobs/{job_id}/anonymize")
        data = result.json()
        for file_result in data.get("files", []):
            if file_result["filename"].endswith(".pdf"):
                assert file_result.get("file_type") == "pdf"
            elif file_result["filename"].endswith(".docx"):
                assert file_result.get("file_type") == "docx"


# ---------------------------------------------------------------------------
# Overlapping PII
# ---------------------------------------------------------------------------

class TestOverlappingPII:
    def test_name_and_email_in_same_sentence_both_replaced(
        self, app_client, sample_pdf_path, mock_llm_endpoint
    ):
        """
        'Contact Jane Smith at jane.smith@acme.com' — both entities replaced.
        We verify this at the service level since controlling exact PDF content
        in the integration test is tricky. Test via the mapper directly.
        """
        from backend.models.schemas import PIIFinding
        from backend.services.mapper import build_mapping
        from backend.services.replacer import apply_replacements

        findings = [
            PIIFinding(text=KNOWN_PII["person"], type="PERSON", confidence="high", source="llm"),
            PIIFinding(text=KNOWN_PII["email"],  type="EMAIL",  confidence="high", source="llm"),
        ]
        mapping = build_mapping(findings)
        text = f"Contact {KNOWN_PII['person']} at {KNOWN_PII['email']}."
        result = apply_replacements(text, mapping)

        assert KNOWN_PII["person"] not in result.text, "Name not replaced"
        assert KNOWN_PII["email"] not in result.text, "Email not replaced"
        assert "[PERSON_1]" in result.text
        assert "[EMAIL_1]" in result.text

    def test_overlapping_spans_treated_as_separate_entities(self):
        """Name substring appearing inside an email must not cause double-replacement."""
        from backend.models.schemas import PIIFinding
        from backend.services.mapper import build_mapping
        from backend.services.replacer import apply_replacements

        # 'jane' appears in the email address too
        findings = [
            PIIFinding(text="Jane Smith",          type="PERSON", confidence="high", source="llm"),
            PIIFinding(text="jane.smith@acme.com", type="EMAIL",  confidence="high", source="llm"),
        ]
        mapping = build_mapping(findings)
        text = "Jane Smith — jane.smith@acme.com"
        result = apply_replacements(text, mapping)

        assert "Jane Smith" not in result.text
        assert "jane.smith@acme.com" not in result.text
        # No partial replacement artifacts like [PERSON_1].smith@acme.com
        assert "@" not in result.text or "[EMAIL_1]" in result.text


# ---------------------------------------------------------------------------
# PII in table cells — DOCX output
# ---------------------------------------------------------------------------

class TestPIIInTableCells:
    def test_table_cell_pii_replaced_in_output(
        self, app_client, sample_docx_path, mock_llm_endpoint
    ):
        job_id, result = anonymize_file(app_client, sample_docx_path, DOCX_MIME)
        data = result.json()
        for file_result in data.get("files", []):
            if file_result["filename"].endswith(".docx"):
                # All text including table cells must have PII replaced
                all_text = file_result.get("anonymized_text", "")
                assert KNOWN_PII["ssn"] not in all_text, (
                    "SSN found in anonymized output — table cell not processed"
                )

    def test_table_structure_preserved_after_replacement(
        self, app_client, sample_docx_path, mock_llm_endpoint
    ):
        """After anonymization, the DOCX table rows/cols must still be intact."""
        import io
        from docx import Document

        job_id, _ = anonymize_file(app_client, sample_docx_path, DOCX_MIME)
        export = app_client.get(f"/jobs/{job_id}/export")
        assert export.status_code == 200

        doc = Document(io.BytesIO(export.content))
        tables = doc.tables
        assert len(tables) >= 1, "No tables in anonymized DOCX"
        # The table should still have 2 rows (header + data)
        assert len(tables[0].rows) == 2
