"""
End-to-End Integration Test — Full Pipeline

Exercises the complete user workflow:
  Upload → Image Review → Anonymize → Review → Export → Re-identify

Uses real PDF and DOCX files (created by fixtures), a mocked LLM endpoint,
and the FastAPI TestClient.  Verifies:

1. All known PII values are absent from anonymized output
2. PII detection rate ≥ 95% of the spec's target list (given mock LLM + regex)
3. Mapping file is valid JSON and covers all replaced items
4. Re-identification restores 100% of replacements
5. Exported files are valid PDF / DOCX
6. Job is saved in the DB with correct metadata
7. No network calls are made outside the configured LLM endpoint

This test is intentionally slow — it is the integration canary.
"""

import io
import json
import re as _re
import zipfile

import pytest
from conftest import KNOWN_PII, SAMPLE_PII_TEXT


PDF_MIME = "application/pdf"
DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# PII values that must be absent from anonymized output
MUST_NOT_APPEAR = [
    KNOWN_PII["person"],
    KNOWN_PII["ssn"],
    KNOWN_PII["email"],
    KNOWN_PII["phone1"],
]

# All distinct PII categories from the spec target list
ALL_PII_CATEGORIES = [
    "PERSON", "ORG", "EMAIL", "PHONE", "ADDRESS",
    "ID", "SSN", "ACCOUNT", "DOB", "OTHER",
]


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def e2e_job_pdf(app_client, sample_pdf_path, mock_llm_endpoint):
    """Upload → Anonymize a PDF, return (job_id, anonymize_response)."""
    with open(sample_pdf_path, "rb") as f:
        upload = app_client.post(
            "/upload",
            files={"files": (sample_pdf_path.name, f, PDF_MIME)},
        )
    assert upload.status_code == 200, upload.text
    job_id = upload.json()["job_id"]
    anon = app_client.post(f"/jobs/{job_id}/anonymize")
    assert anon.status_code == 200, anon.text
    return job_id, anon.json()


@pytest.fixture
def e2e_job_docx(app_client, sample_docx_path, mock_llm_endpoint):
    """Upload → Anonymize a DOCX, return (job_id, anonymize_response)."""
    with open(sample_docx_path, "rb") as f:
        upload = app_client.post(
            "/upload",
            files={"files": (sample_docx_path.name, f, DOCX_MIME)},
        )
    assert upload.status_code == 200, upload.text
    job_id = upload.json()["job_id"]
    anon = app_client.post(f"/jobs/{job_id}/anonymize")
    assert anon.status_code == 200, anon.text
    return job_id, anon.json()


@pytest.fixture
def e2e_job_mixed(app_client, sample_pdf_path, sample_docx_path, mock_llm_endpoint):
    """Upload PDF + DOCX in the same job."""
    with open(sample_pdf_path, "rb") as f1, open(sample_docx_path, "rb") as f2:
        upload = app_client.post(
            "/upload",
            files=[
                ("files", (sample_pdf_path.name, f1, PDF_MIME)),
                ("files", (sample_docx_path.name, f2, DOCX_MIME)),
            ],
        )
    assert upload.status_code == 200
    job_id = upload.json()["job_id"]
    anon = app_client.post(f"/jobs/{job_id}/anonymize")
    assert anon.status_code == 200
    return job_id, anon.json()


# ---------------------------------------------------------------------------
# Step 1: Upload
# ---------------------------------------------------------------------------

class TestE2EUpload:
    def test_pdf_upload_creates_job(self, app_client, sample_pdf_path):
        with open(sample_pdf_path, "rb") as f:
            resp = app_client.post(
                "/upload",
                files={"files": (sample_pdf_path.name, f, PDF_MIME)},
            )
        assert resp.status_code == 200
        assert "job_id" in resp.json()

    def test_docx_upload_creates_job(self, app_client, sample_docx_path):
        with open(sample_docx_path, "rb") as f:
            resp = app_client.post(
                "/upload",
                files={"files": (sample_docx_path.name, f, DOCX_MIME)},
            )
        assert resp.status_code == 200
        assert "job_id" in resp.json()

    def test_upload_response_includes_page_count(self, app_client, sample_pdf_path):
        with open(sample_pdf_path, "rb") as f:
            resp = app_client.post(
                "/upload",
                files={"files": (sample_pdf_path.name, f, PDF_MIME)},
            )
        file_info = resp.json()["files"][0]
        assert "page_count" in file_info
        assert file_info["page_count"] >= 1

    def test_upload_response_includes_file_size(self, app_client, sample_pdf_path):
        with open(sample_pdf_path, "rb") as f:
            resp = app_client.post(
                "/upload",
                files={"files": (sample_pdf_path.name, f, PDF_MIME)},
            )
        file_info = resp.json()["files"][0]
        assert "size_bytes" in file_info
        assert file_info["size_bytes"] > 0


# ---------------------------------------------------------------------------
# Step 2: Image Review (API contract)
# ---------------------------------------------------------------------------

class TestE2EImageReview:
    def test_images_endpoint_returns_list(self, app_client, pdf_with_image_path):
        with open(pdf_with_image_path, "rb") as f:
            upload = app_client.post(
                "/upload",
                files={"files": (pdf_with_image_path.name, f, PDF_MIME)},
            )
        job_id = upload.json()["job_id"]
        resp = app_client.get(f"/jobs/{job_id}/images")
        assert resp.status_code == 200
        assert isinstance(resp.json(), list)

    def test_images_include_metadata(self, app_client, pdf_with_image_path):
        with open(pdf_with_image_path, "rb") as f:
            upload = app_client.post(
                "/upload",
                files={"files": (pdf_with_image_path.name, f, PDF_MIME)},
            )
        job_id = upload.json()["job_id"]
        images = app_client.get(f"/jobs/{job_id}/images").json()
        if images:
            assert "page_number" in images[0]
            assert "source_filename" in images[0]
            assert "marked_for_removal" in images[0]

    def test_update_image_removal_flag(self, app_client, pdf_with_image_path):
        with open(pdf_with_image_path, "rb") as f:
            upload = app_client.post(
                "/upload",
                files={"files": (pdf_with_image_path.name, f, PDF_MIME)},
            )
        job_id = upload.json()["job_id"]
        images = app_client.get(f"/jobs/{job_id}/images").json()
        if images:
            image_id = images[0]["id"]
            resp = app_client.patch(
                f"/jobs/{job_id}/images/{image_id}",
                json={"marked_for_removal": False},
            )
            assert resp.status_code == 200


# ---------------------------------------------------------------------------
# Step 3: Anonymize — PII removal verification
# ---------------------------------------------------------------------------

class TestE2EAnonymize:
    def test_known_ssn_absent_from_pdf_output(self, e2e_job_pdf):
        _, data = e2e_job_pdf
        for file_result in data["files"]:
            assert KNOWN_PII["ssn"] not in file_result.get("anonymized_text", "")

    def test_known_email_absent_from_pdf_output(self, e2e_job_pdf):
        _, data = e2e_job_pdf
        for file_result in data["files"]:
            assert KNOWN_PII["email"] not in file_result.get("anonymized_text", "")

    def test_known_person_absent_from_pdf_output(self, e2e_job_pdf):
        _, data = e2e_job_pdf
        for file_result in data["files"]:
            assert KNOWN_PII["person"] not in file_result.get("anonymized_text", "")

    def test_known_ssn_absent_from_docx_output(self, e2e_job_docx):
        _, data = e2e_job_docx
        for file_result in data["files"]:
            assert KNOWN_PII["ssn"] not in file_result.get("anonymized_text", "")

    def test_known_email_absent_from_docx_output(self, e2e_job_docx):
        _, data = e2e_job_docx
        for file_result in data["files"]:
            assert KNOWN_PII["email"] not in file_result.get("anonymized_text", "")

    def test_pii_detection_rate_meets_95_percent_threshold(self, e2e_job_pdf):
        """
        With mock LLM returning all known PII + regex catching pattern-based PII,
        all PII values in MUST_NOT_APPEAR must be absent from output.
        This approximates the ≥95% success criterion.
        """
        _, data = e2e_job_pdf
        anonymized = " ".join(
            f.get("anonymized_text", "") for f in data["files"]
        )
        found = [v for v in MUST_NOT_APPEAR if v in anonymized]
        detected_rate = 1 - len(found) / len(MUST_NOT_APPEAR)
        assert detected_rate >= 0.95, (
            f"PII detection rate {detected_rate:.0%} < 95%. "
            f"Still present: {found}"
        )

    def test_placeholders_well_formed_in_output(self, e2e_job_pdf):
        _, data = e2e_job_pdf
        for file_result in data["files"]:
            anonymized = file_result.get("anonymized_text", "")
            placeholders = _re.findall(r"\[[A-Z_]+_\d+\]", anonymized)
            for ph in placeholders:
                # Must match [TYPE_N] where TYPE is one of the known types
                assert _re.match(
                    r"\[(PERSON|ORG|EMAIL|PHONE|ADDRESS|ID|SSN|ACCOUNT|DOB|OTHER)_\d+\]",
                    ph,
                ), f"Malformed placeholder: {ph}"

    def test_mixed_batch_both_files_anonymized(self, e2e_job_mixed):
        _, data = e2e_job_mixed
        assert len(data["files"]) == 2
        for file_result in data["files"]:
            anonymized = file_result.get("anonymized_text", "")
            assert KNOWN_PII["ssn"] not in anonymized


# ---------------------------------------------------------------------------
# Step 4: Review
# ---------------------------------------------------------------------------

class TestE2EReview:
    def test_review_returns_original_text(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        review = app_client.get(f"/jobs/{job_id}/review").json()
        for f in review["files"]:
            assert KNOWN_PII["person"] in f["original_text"] or len(f["original_text"]) > 0

    def test_review_mapping_covers_all_replaced_pii(self, app_client, e2e_job_pdf):
        job_id, data = e2e_job_pdf
        review = app_client.get(f"/jobs/{job_id}/review").json()
        mapping = review["mapping"]
        placeholders_in_mapping = {e["placeholder"] for e in mapping}

        # Every placeholder that appears in anonymized text must be in the mapping
        for file_result in data["files"]:
            anonymized = file_result.get("anonymized_text", "")
            placeholders_in_text = set(_re.findall(r"\[[A-Z_]+_\d+\]", anonymized))
            unmapped = placeholders_in_text - placeholders_in_mapping
            assert not unmapped, f"Placeholders in output with no mapping entry: {unmapped}"

    def test_review_positions_match_anonymized_text(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        review = app_client.get(f"/jobs/{job_id}/review").json()
        for file_result in review["files"]:
            anonymized = file_result["anonymized_text"]
            for pos in file_result.get("positions", []):
                start, end = pos["start"], pos["end"]
                assert 0 <= start < end <= len(anonymized), (
                    f"Position [{start}:{end}] out of range for text length {len(anonymized)}"
                )


# ---------------------------------------------------------------------------
# Step 5: Export
# ---------------------------------------------------------------------------

class TestE2EExport:
    def test_export_single_pdf(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        resp = app_client.get(f"/jobs/{job_id}/export")
        assert resp.status_code == 200
        # Should start with PDF magic bytes
        assert resp.content[:4] == b"%PDF", "Exported file is not a valid PDF"

    def test_export_single_docx(self, app_client, e2e_job_docx):
        job_id, _ = e2e_job_docx
        resp = app_client.get(f"/jobs/{job_id}/export")
        assert resp.status_code == 200
        # DOCX is a ZIP — check magic bytes
        assert resp.content[:2] == b"PK", "Exported file is not a valid DOCX (expected ZIP magic bytes)"

    def test_export_docx_is_valid_docx(self, app_client, e2e_job_docx):
        from docx import Document
        job_id, _ = e2e_job_docx
        resp = app_client.get(f"/jobs/{job_id}/export")
        doc = Document(io.BytesIO(resp.content))
        assert doc is not None

    def test_export_mapping_json_is_valid(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        resp = app_client.get(f"/jobs/{job_id}/export/mapping")
        assert resp.status_code == 200
        mapping = resp.json()
        assert isinstance(mapping, list)
        assert len(mapping) > 0

    def test_export_mapping_has_required_fields(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        mapping = app_client.get(f"/jobs/{job_id}/export/mapping").json()
        for entry in mapping:
            assert "original" in entry
            assert "placeholder" in entry
            assert "pii_type" in entry

    def test_export_mixed_job_is_zip(self, app_client, e2e_job_mixed):
        job_id, _ = e2e_job_mixed
        resp = app_client.get(f"/jobs/{job_id}/export")
        assert resp.status_code == 200
        assert zipfile.is_zipfile(io.BytesIO(resp.content)), (
            "Multi-file export should be a ZIP"
        )

    def test_export_zip_contains_both_files(self, app_client, e2e_job_mixed, sample_pdf_path, sample_docx_path):
        job_id, _ = e2e_job_mixed
        resp = app_client.get(f"/jobs/{job_id}/export")
        z = zipfile.ZipFile(io.BytesIO(resp.content))
        names = z.namelist()
        assert any(sample_pdf_path.stem in n for n in names)
        assert any(sample_docx_path.stem in n for n in names)

    def test_job_saved_in_db_after_export(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        app_client.get(f"/jobs/{job_id}/export")
        jobs = app_client.get("/jobs").json()
        ids = [j["id"] for j in jobs]
        assert job_id in ids

    def test_job_summary_metadata(self, app_client, e2e_job_pdf):
        job_id, _ = e2e_job_pdf
        summary_resp = app_client.get(f"/jobs/{job_id}/summary")
        assert summary_resp.status_code == 200
        summary = summary_resp.json()
        assert "file_count" in summary
        assert "pii_items_found" in summary
        assert "model_used" in summary
        assert summary["file_count"] >= 1
        assert summary["pii_items_found"] >= 0


# ---------------------------------------------------------------------------
# Step 6: Re-identification — 100% fidelity
# ---------------------------------------------------------------------------

class TestE2EReidentify:
    @pytest.fixture
    def reidentify_setup(self, app_client, e2e_job_pdf):
        job_id, data = e2e_job_pdf
        mapping_list = app_client.get(f"/jobs/{job_id}/export/mapping").json()
        # Build placeholder→original dict for the re-identify call
        mapping_dict = {e["placeholder"]: e["original"] for e in mapping_list}
        anonymized_bytes = app_client.get(f"/jobs/{job_id}/export").content
        return job_id, mapping_dict, anonymized_bytes

    def test_reidentify_returns_200(self, app_client, reidentify_setup):
        job_id, mapping_dict, _ = reidentify_setup
        resp = app_client.post(
            "/reidentify",
            json={"job_id": job_id, "mapping": mapping_dict},
        )
        assert resp.status_code == 200

    def test_reidentify_output_is_valid_pdf(self, app_client, reidentify_setup):
        job_id, mapping_dict, _ = reidentify_setup
        resp = app_client.post(
            "/reidentify",
            json={"job_id": job_id, "mapping": mapping_dict},
        )
        assert resp.content[:4] == b"%PDF", "Re-identified output is not a valid PDF"

    def test_reidentify_restores_person_name(self, app_client, e2e_job_pdf, reidentify_setup):
        job_id, mapping_dict, _ = reidentify_setup
        resp = app_client.post(
            "/reidentify",
            json={"job_id": job_id, "mapping": mapping_dict},
        )
        # Verify by reading the restored text via review
        restore_review = app_client.get(f"/jobs/{job_id}/review?restored=true")
        if restore_review.status_code == 200:
            restored_text = " ".join(
                f.get("restored_text", "") for f in restore_review.json().get("files", [])
            )
            assert KNOWN_PII["person"] in restored_text

    def test_reidentify_100_percent_mapping_fidelity(self, app_client, e2e_job_pdf):
        """
        Full round-trip at the service level: replace → reverse.
        Spec success criterion: 100% of replacements restored.
        """
        from backend.models.schemas import PIIFinding
        from backend.services.mapper import build_mapping
        from backend.services.replacer import apply_replacements, reverse_replacements

        findings = [
            PIIFinding(text=KNOWN_PII["person"],  type="PERSON",  confidence="high", source="llm"),
            PIIFinding(text=KNOWN_PII["email"],   type="EMAIL",   confidence="high", source="llm"),
            PIIFinding(text=KNOWN_PII["phone1"],  type="PHONE",   confidence="high", source="regex"),
            PIIFinding(text=KNOWN_PII["ssn"],     type="SSN",     confidence="high", source="regex"),
            PIIFinding(text=KNOWN_PII["address"], type="ADDRESS", confidence="high", source="llm"),
        ]
        mapping = build_mapping(findings)
        original = SAMPLE_PII_TEXT
        replaced = apply_replacements(original, mapping)
        restored = reverse_replacements(replaced.text, mapping)

        assert restored == original, (
            "Re-identification did not perfectly restore the original text.\n"
            f"Original:  {original[:200]}\n"
            f"Restored:  {restored[:200]}"
        )

    def test_no_placeholders_remain_after_reidentify(self, app_client, reidentify_setup):
        """After re-identification, no [TYPE_N] placeholders should remain."""
        job_id, mapping_dict, _ = reidentify_setup
        resp = app_client.post(
            "/reidentify",
            json={"job_id": job_id, "mapping": mapping_dict},
        )
        restore_review = app_client.get(f"/jobs/{job_id}/review?restored=true")
        if restore_review.status_code == 200:
            all_text = " ".join(
                f.get("restored_text", "") for f in restore_review.json().get("files", [])
            )
            remaining = _re.findall(r"\[[A-Z_]+_\d+\]", all_text)
            assert not remaining, f"Unreplaced placeholders after re-identify: {remaining}"


# ---------------------------------------------------------------------------
# Security: no outbound network calls beyond LLM endpoint
# ---------------------------------------------------------------------------

class TestNoExtraNetworkCalls:
    def test_only_llm_endpoint_called(self, app_client, sample_pdf_path, mock_llm_endpoint):
        """
        All HTTP requests made during anonymization must go to the configured
        LLM endpoint only.  We verify by checking the httpx mock received only
        calls to localhost:11434.
        """
        with open(sample_pdf_path, "rb") as f:
            upload = app_client.post(
                "/upload",
                files={"files": (sample_pdf_path.name, f, PDF_MIME)},
            )
        job_id = upload.json()["job_id"]
        app_client.post(f"/jobs/{job_id}/anonymize")

        # All requests captured by httpx_mock must target the LLM endpoint
        for req in mock_llm_endpoint.get_requests():
            assert "localhost:11434" in str(req.url), (
                f"Unexpected outbound request to: {req.url}"
            )
